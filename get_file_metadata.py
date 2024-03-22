import os
import PyPDF2
from openpyxl import load_workbook
from docx import Document

def extract_docx_metadata(docx_file):       #Extraer los metadatos de un docx
    doc = Document(docx_file)
    metadata = {}
    
    # Extract document properties
    properties = doc.core_properties
    metadata['Title'] = properties.title
    metadata['Author'] = properties.author
    metadata['Subject'] = properties.subject
    metadata['Keywords'] = properties.keywords
    metadata['Comments'] = properties.comments
    
    return metadata

def extract_xlsx_metadata(xlsx_file):       #Extraer los metadatos de un xlsx
    wb = load_workbook(filename=xlsx_file)
    metadata = {}
    
    # Extraer propiedades del libro de trabajo
    props = wb.properties
    metadata['Title'] = props.title
    metadata['Author'] = props.creator
    metadata['Last Modified By'] = props.lastModifiedBy
    metadata['Created'] = props.created
    metadata['Modified'] = props.modified
    
    return metadata

def extract_metadata(pdf_file):         #Extraer los metadatos de un PDF
    with open(pdf_file, 'rb') as f:
        reader = PyPDF2.PdfFileReader(f)
        metadata = reader.getDocumentInfo()
    return metadata

def main(folder_path):
    # Verificar si la carpeta existe
    if not os.path.isdir(folder_path):
        print("La carpeta especificada no existe.")
        return

    # Recorrer todos los archivos en la carpeta
    for filename in os.listdir(folder_path):
        if filename.endswith('.docx'):
            file_path = os.path.join(folder_path, filename)
            metadata = extract_docx_metadata(file_path)
            print(f"Metadatos de '{filename}':")
            for key, value in metadata.items():
                print(f"{key}: {value}")
            print()
        elif filename.endswith('.xlsx'):
            file_path = os.path.join(folder_path, filename)
            metadata = extract_xlsx_metadata(file_path)
            print(f"Metadatos de '{filename}':")
            for key, value in metadata.items():
                print(f"{key}: {value}")
            print()
        elif filename.endswith('.pdf'):
            file_path = os.path.join(folder_path, filename)
            metadata = extract_metadata(file_path)
            print(f"Metadatos de '{filename}':")
            for key, value in metadata.items():
                print(f"{key}: {value}")
            print()

if __name__ == "__main__":
    folder_path = input("Ingrese la ruta de la carpeta que contiene los archivos: ")
    main(folder_path)

